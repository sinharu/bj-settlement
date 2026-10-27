from io import BytesIO

import pytest
from openpyxl import load_workbook

from heart_text import (
    make_heart_excel,
    parse_heart_text,
    sheet_rows,
    total_mismatch,
)


@pytest.fixture
def heart_text():
    return "\n".join([
        "테스트방송A   1,250",
        "user01(닉네임A)   200",
        "00002(닉네임B)\t1,000",
        "user03(닉네임C)   50",
        "총합계   1,250",
    ])


def test_basic_parsing_and_descending_sort(heart_text):
    result = parse_heart_text(heart_text)

    assert result.name == "테스트방송A"
    assert result.name_total == 1250
    assert result.total == 1250
    assert result.rows == [
        {"아이디": "00002", "닉네임": "닉네임B", "후원하트": 1000},
        {"아이디": "user01", "닉네임": "닉네임A", "후원하트": 200},
        {"아이디": "user03", "닉네임": "닉네임C", "후원하트": 50},
    ]
    assert result.sum == 1250
    assert result.skipped == []


def test_sort_ties_keep_input_order():
    result = parse_heart_text("\n".join([
        "user03(닉네임C)   20",
        "user02(닉네임B)   50",
        "user01(닉네임A)   20",
        "user05(닉네임E)   50",
        "user04(닉네임D)   20",
    ]))

    assert [row["아이디"] for row in result.rows] == [
        "user02", "user05", "user03", "user01", "user04",
    ]
    assert [row["후원하트"] for row in result.rows] == [50, 50, 20, 20, 20]


@pytest.mark.parametrize("nickname", [
    "닉네임A(별칭B)",
    "[닉네임A]😀™",
    "닉네임A（별칭B）【테스트】",
    "닉네임A(별칭B)[태그C]💖™",
    "닉네임A)별칭B(태그C",
    "닉네임 A [태그 B] 👩‍💻",
])
def test_nickname_uses_first_open_and_last_close_parenthesis(nickname):
    result = parse_heart_text(f"user01({nickname})   10")

    assert result.rows == [
        {"아이디": "user01", "닉네임": nickname, "후원하트": 10},
    ]
    assert result.skipped == []


@pytest.mark.parametrize("separator", [" ", "   ", "\t", " \t "])
def test_whitespace_and_comma_numbers(separator):
    result = parse_heart_text("\n".join([
        f"  테스트 방송A{separator}1,234,567  ",
        f"  user01(닉네임A){separator}1,234,567  ",
        f"  총합계{separator}1,234,567  ",
    ]))

    assert result.name == "테스트 방송A"
    assert result.name_total == 1234567
    assert result.total == 1234567
    assert result.sum == 1234567
    assert result.rows[0]["후원하트"] == 1234567
    assert result.skipped == []


def test_sheet_rows_layout_and_numeric_id(heart_text):
    result = parse_heart_text(heart_text)

    assert sheet_rows(result) == [
        ["테스트방송A", None, 1250],
        ["아이디", "닉네임", "후원하트"],
        ["00002", "닉네임B", 1000],
        ["user01", "닉네임A", 200],
        ["user03", "닉네임C", 50],
        ["총합계", None, 1250],
    ]
    assert result.rows[0]["아이디"] == "00002"


@pytest.mark.parametrize(("user_id", "expected"), [
    ("0", 0),
    ("00002", "00002"),
    ("4742795197", 4742795197),
    ("111111111111111", 111111111111111),
    ("user01", "user01"),
    ("01user", "01user"),
])
def test_digit_only_ids_become_integer_cells_without_losing_zeros(user_id, expected):
    result = parse_heart_text(f"{user_id}(닉네임A)   10")
    cell_value = sheet_rows(result)[2][0]

    assert cell_value == expected
    assert type(cell_value) is type(expected)
    assert result.rows[0]["아이디"] == user_id


def test_total_line_is_not_a_donor_or_streamer():
    result = parse_heart_text("총합계   30\nuser01(닉네임A)   30")

    assert result.name == ""
    assert result.name_total is None
    assert result.total == 30
    assert result.rows == [
        {"아이디": "user01", "닉네임": "닉네임A", "후원하트": 30},
    ]
    assert result.skipped == []


@pytest.mark.parametrize(("name_total", "stated_total", "expected"), [
    (30, 30, None),
    (30, 35, 35),
    (35, 30, None),
    (35, 40, 40),
    (30, None, None),
    (35, None, 35),
    (None, 30, None),
    (None, 35, 35),
    (None, None, None),
    (30, 0, 0),
    (0, None, 0),
])
def test_total_mismatch_uses_stated_total_then_streamer_total(
    name_total, stated_total, expected,
):
    lines = []
    if name_total is not None:
        lines.append(f"테스트방송A   {name_total}")
    lines.extend(["user01(닉네임A)   10", "user02(닉네임B)   20"])
    if stated_total is not None:
        lines.append(f"총합계   {stated_total}")

    result = parse_heart_text("\n".join(lines))

    assert result.sum == 30
    assert total_mismatch(result) == expected
    if stated_total is not None:
        assert sheet_rows(result)[-1] == ["총합계", None, stated_total]


def test_blank_lines_ignored_and_unparsable_lines_recorded_in_order():
    invalid_lines = [
        "하트 없는 줄",
        "user02(닉네임B) 숫자아님",
        "user03(닉네임C) 1.5",
        "user04(닉네임D 10",
        "user05닉네임E) 10",
        "이름없는줄 10",
    ]
    result = parse_heart_text("\n".join([
        "", " \t ", "user01(닉네임A)   10", "",
        *[f"  {line}  " for line in invalid_lines],
        "\t", "user06(닉네임F)   20", "",
    ]))

    assert result.skipped == invalid_lines
    assert [row["아이디"] for row in result.rows] == ["user06", "user01"]
    assert result.sum == 30


def test_first_valid_name_line_before_donors_is_used_once():
    result = parse_heart_text("\n".join([
        "", "숫자 없는 줄", "테스트방송A   30", "테스트방송B   999",
        "user01(닉네임A)   30",
    ]))

    assert result.name == "테스트방송A"
    assert result.name_total == 30
    assert result.skipped == ["숫자 없는 줄", "테스트방송B   999"]
    assert result.sum == 30


@pytest.mark.parametrize("total_line", ["", "\n총합계   30"])
def test_missing_first_line_preserves_donors_and_leaves_name_unset(total_line):
    result = parse_heart_text(
        "user01(닉네임A)   10\nuser02(닉네임B)   20" + total_line,
    )

    assert result.name == ""
    assert result.name_total is None
    assert result.total == (30 if total_line else None)
    assert result.rows == [
        {"아이디": "user02", "닉네임": "닉네임B", "후원하트": 20},
        {"아이디": "user01", "닉네임": "닉네임A", "후원하트": 10},
    ]
    assert result.skipped == []
    assert total_mismatch(result) is None


@pytest.mark.parametrize("text", ["", " \n\t\r\n "])
def test_empty_input(text):
    result = parse_heart_text(text)

    assert result.name == ""
    assert result.name_total is None
    assert result.total is None
    assert result.rows == []
    assert result.skipped == []
    assert result.sum == 0
    assert total_mismatch(result) is None


def test_zero_hearts_are_preserved():
    result = parse_heart_text(
        "테스트방송A   0\nuser01(닉네임A)   0\n총합계   0",
    )

    assert sheet_rows(result) == [
        ["테스트방송A", None, 0],
        ["아이디", "닉네임", "후원하트"],
        ["user01", "닉네임A", 0],
        ["총합계", None, 0],
    ]
    assert total_mismatch(result) is None


def test_make_heart_excel_round_trip_values_and_formatting(heart_text):
    result = parse_heart_text(heart_text)
    excel = make_heart_excel(result)

    assert isinstance(excel, BytesIO)
    assert excel.tell() == 0
    workbook = load_workbook(excel)
    try:
        assert workbook.sheetnames == ["Sheet1"]
        sheet = workbook["Sheet1"]
        assert sheet.max_row == 6
        assert sheet.max_column == 3
        assert list(sheet.iter_rows(values_only=True)) == [
            ("테스트방송A", None, 1250),
            ("아이디", "닉네임", "후원하트"),
            ("00002", "닉네임B", 1000),
            ("user01", "닉네임A", 200),
            ("user03", "닉네임C", 50),
            ("총합계", None, 1250),
        ]
        assert sheet["A3"].data_type == "s"
        assert sheet["A4"].data_type == "s"
        for row in sheet.iter_rows(min_row=1, max_row=6, min_col=1, max_col=3):
            for cell in row:
                for side in ("left", "right", "top", "bottom"):
                    assert getattr(cell.border, side).style == "thin", cell.coordinate
                assert cell.font.name == "맑은 고딕", cell.coordinate
                assert cell.font.sz == 11, cell.coordinate
        for cell in sheet["C"]:
            assert cell.number_format == r'_-* #,##0_-;\-* #,##0_-;_-* "-"_-;_-@_-'
        for column, width in {"A": 29.5, "B": 26.0, "C": 9.375}.items():
            assert sheet.column_dimensions[column].width == pytest.approx(width)
    finally:
        workbook.close()
        excel.close()


def test_make_heart_excel_keeps_formula_like_text_as_text():
    result = parse_heart_text("테스트방송A   10\nuser01(=1+1)   10")
    workbook = load_workbook(make_heart_excel(result))
    try:
        cell = workbook["Sheet1"]["B3"]
        assert cell.value == "=1+1"
        assert cell.data_type == "s"
    finally:
        workbook.close()
